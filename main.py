import json
import csv
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment
from src import utils
from src import indices

# ----------------------------
# Configuration settings
# ----------------------------
json_file = "./file_to_process/content.json"
MAX_WORKS = None  # Set to an integer to limit processing, or None to process all

# ----------------------------
# Index registry
# ----------------------------

# Maps index name → callable(text, lang) → score
INDEX_REGISTRY = {
    "gulpease":    lambda text, lang: indices.gulpease_index(text),
    "flesch":      lambda text, lang: indices.flesch_index(text, lang),
    "gunning_fog": lambda text, lang: indices.gunning_fog_index(text, lang),
}

# Maps index name → set of languages it supports
INDEX_LANGS = {
    "gulpease":    {"it"},
    "flesch":      {"it", "en"},
    "gunning_fog": {"it", "en"},
}

# ----------------------------
# Load JSON data from file
# ----------------------------
with open(json_file, "r", encoding="utf-8") as f:
    data = json.load(f)

# ----------------------------
# Shared data collection
# ----------------------------

def _collect_results(indices_to_use, lang, process_all_categories, category, sub_category):
    """
    Collect readability metrics for each opera in the corpus.

    Returns:
        results       : list of dicts, one per opera, with all computed columns
        indices_to_use: resolved list of index names (in case "all" was passed)
        lang          : original lang argument (for filename generation)
        cat_label     : "all" or "{category}_{sub_category}" (for filename generation)
    """
    if indices_to_use == "all":
        indices_to_use = list(INDEX_REGISTRY.keys())

    langs_to_process = list(data.keys()) if lang == "all" else [lang]

    tagged_works = []
    for current_lang in langs_to_process:
        if current_lang not in data:
            raise ValueError(f"Language '{current_lang}' not found in the JSON data.")
        if process_all_categories:
            for age_group in data[current_lang]:
                for cat in data[current_lang][age_group]:
                    for title, sentences in utils.extract_works(data[current_lang][age_group][cat]):
                        tagged_works.append((title, " ".join(sentences), "\n".join(sentences), current_lang))
        else:
            if category is None or sub_category is None:
                raise ValueError("Both 'category' and 'sub_category' must be specified if 'process_all_categories' is False.")
            for title, sentences in utils.extract_works(data[current_lang][category][sub_category]):
                tagged_works.append((title, " ".join(sentences), "\n".join(sentences), current_lang))

    if MAX_WORKS:
        tagged_works = tagged_works[:MAX_WORKS]

    results = []
    for idx, (title, analysis_text, display_text, current_lang) in enumerate(tagged_works, start=1):
        row = {
            "Id": idx,
            "Lang": current_lang,
            "Title": title,
            "Sentence_Text": display_text,
            "Letters": utils.letter_count(analysis_text),
            "Words": utils.word_count(analysis_text),
            "Num_Sentences": utils.sentence_count(analysis_text),
        }
        for index_name in indices_to_use:
            if current_lang in INDEX_LANGS[index_name]:
                score = INDEX_REGISTRY[index_name](analysis_text, current_lang)
                row[index_name] = round(score, 2) if score is not None else None
            else:
                row[index_name] = "N/A"
        results.append(row)

    cat_label = "all" if process_all_categories else f"{category}_{sub_category}"
    return results, indices_to_use, lang, cat_label

# ----------------------------
# Report generation
# ----------------------------

def generate_csv_report(indices_to_use, lang, process_all_categories=False, category=None, sub_category=None):
    """
    Generate a CSV readability report.

    Parameters:
        indices_to_use : list[str] | "all"
        lang           : str | "all"
        process_all_categories : bool
        category       : str | None
        sub_category   : str | None
    """
    results, indices_to_use, lang, cat_label = _collect_results(
        indices_to_use, lang, process_all_categories, category, sub_category
    )

    os.makedirs("./reports", exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    indices_label = "_".join(indices_to_use)
    path = f"./reports/report_{indices_label}_{lang}_{cat_label}_{timestamp}.csv"

    fieldnames = ["Id", "Lang", "Title", "Sentence_Text", "Letters", "Words", "Num_Sentences"] + indices_to_use
    with open(path, "w", newline="", encoding="utf-8") as csvfile:
        writer = csv.DictWriter(csvfile, fieldnames=fieldnames, quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        for row in results:
            writer.writerow(row)

    print(f"CSV report generated: {path}")


def generate_excel_report(indices_to_use, lang, process_all_categories=False, category=None, sub_category=None):
    """
    Generate an Excel (.xlsx) readability report identical in content to the CSV,
    with bold header row, auto-sized columns, and text-wrap on the Sentence_Text column.

    Parameters:
        indices_to_use : list[str] | "all"
        lang           : str | "all"
        process_all_categories : bool
        category       : str | None
        sub_category   : str | None
    """
    results, indices_to_use, lang, cat_label = _collect_results(
        indices_to_use, lang, process_all_categories, category, sub_category
    )

    fieldnames = ["Id", "Lang", "Title", "Sentence_Text", "Letters", "Words", "Num_Sentences"] + indices_to_use

    wb = openpyxl.Workbook()
    ws = wb.active

    # Header row — bold
    bold = Font(bold=True)
    for col_idx, name in enumerate(fieldnames, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = bold

    wrap = Alignment(wrap_text=True, vertical="top")
    no_wrap = Alignment(vertical="top")
    sentence_col = fieldnames.index("Sentence_Text") + 1  # 1-based column index

    # Data rows
    for row_idx, row in enumerate(results, start=2):
        for col_idx, name in enumerate(fieldnames, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[name])
            cell.alignment = wrap if col_idx == sentence_col else no_wrap

    # Auto-size column widths (capped at 60)
    for col_idx, name in enumerate(fieldnames, start=1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = len(name)
        for row in results:
            val = str(row[name])
            line_max = max(len(line) for line in val.splitlines()) if "\n" in val else len(val)
            max_len = max(max_len, line_max)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

    os.makedirs("./reports", exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    indices_label = "_".join(indices_to_use)
    path = f"./reports/report_{indices_label}_{lang}_{cat_label}_{timestamp}.xlsx"
    wb.save(path)

    print(f"Excel report generated: {path}")


# ----------------------------
# Examples
# ----------------------------

# generate_csv_report(indices_to_use=["gulpease"], lang="it", process_all_categories=False, category="adult", sub_category="typical")
# generate_csv_report(indices_to_use="all", lang="all", process_all_categories=True)
# generate_excel_report(indices_to_use="all", lang="all", process_all_categories=True)
generate_csv_report(indices_to_use="all", lang="it", process_all_categories=False, category="adult", sub_category="typical")
generate_excel_report(indices_to_use="all", lang="it", process_all_categories=False, category="adult", sub_category="typical")
